"""Compact marker-band calibration from concentration/Mean/std XLSX workbooks."""
from __future__ import annotations

import os
from dataclasses import dataclass

import numpy as np
from scipy.optimize import curve_fit


@dataclass
class SummaryCurve:
    name: str
    concentration_M: np.ndarray
    mean: np.ndarray
    std: np.ndarray
    offset: float = 0.0
    plateau: float = 0.0
    K: float = 0.0
    m: float = 1.0

    def predict(self, concentration_M):
        c = np.asarray(concentration_M, float)
        u = np.power(np.clip(self.K * c, 0.0, None), self.m)
        return self.offset + self.plateau * u / (1.0 + u)

    def invert(self, response):
        y = np.asarray(response, float) - self.offset
        frac = np.clip(y / max(self.plateau, 1e-12), 0.0, 1.0 - 1e-6)
        u = frac / np.maximum(1.0 - frac, 1e-12)
        return np.power(u, 1.0 / max(self.m, 1e-6)) / max(self.K, 1e-30)


def _fit(curve):
    C, B, sd = curve.concentration_M, curve.mean, curve.std
    b0 = max(0.0, float(np.min(B)) * 0.9)
    p0 = max(float(np.max(B) - b0) * 1.2, 1e-9)
    k0 = 1.0 / max(float(np.median(C)), 1e-12)

    def model(c, offset, plateau, K, m):
        u = np.power(np.clip(K * c, 0.0, None), m)
        return offset + plateau * u / (1.0 + u)

    sigma = np.maximum(np.asarray(sd, float), np.nanmedian(sd) * 0.1 + 1e-9)
    try:
        popt, _ = curve_fit(model, C, B, p0=[b0, p0, k0, 1.0], sigma=sigma,
                            absolute_sigma=False,
                            bounds=([0.0, 0.0, 0.0, 0.05],
                                    [np.inf, np.inf, np.inf, 3.0]), maxfev=50000)
        curve.offset, curve.plateau, curve.K, curve.m = map(float, popt)
    except Exception:
        curve.offset, curve.plateau, curve.K, curve.m = b0, p0, k0, 1.0
    return curve


def load_summary_calibration(path):
    """Return (curves, corrections) from two-row concentration/Mean/std blocks."""
    if os.path.splitext(path)[1].lower() not in (".xlsx", ".xlsm"):
        raise ValueError("summary calibration must be an .xlsx workbook")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("reading summary calibration needs openpyxl") from exc
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    blocks = []
    for col in range(2, ws.max_column + 1):
        name, mh = ws.cell(1, col).value, ws.cell(2, col).value
        sh = ws.cell(2, col + 1).value if col < ws.max_column else None
        if (isinstance(name, str) and name.strip()
                and str(mh).strip().lower() == "mean"
                and str(sh).strip().lower() in ("std", "sd")):
            blocks.append((name.strip(), col - 1, col, col + 1))
    if not blocks:
        raise ValueError("no compound concentration/Mean/std blocks found")
    raw = []
    for name, cc, mc, sc in blocks:
        rows = []
        for row in range(3, ws.max_row + 1):
            vals = (ws.cell(row, cc).value, ws.cell(row, mc).value, ws.cell(row, sc).value)
            if all(v is None for v in vals):
                continue
            try:
                rows.append([float(v) for v in vals])
            except (TypeError, ValueError) as exc:
                raise ValueError(f"{name} row {row} is not numeric") from exc
        if len(rows) < 4:
            raise ValueError(f"{name} needs at least four calibration levels")
        raw.append([name, rows])
    wb.close()

    corrections = []
    if len({len(rows) for _name, rows in raw}) == 1:
        for i in range(len(raw[0][1])):
            cs = np.array([rows[i][0] for _name, rows in raw])
            vals, counts = np.unique(cs, return_counts=True)
            target = float(vals[int(np.argmax(counts))])
            if counts.max() >= 2:
                for name, rows in raw:
                    old = rows[i][0]
                    if old != target and abs(old-target) <= 0.20*max(target, 1e-12):
                        rows[i][0] = target
                        corrections.append(f"{name}: {old:g} -> {target:g} uM")

    curves = []
    for name, rows in raw:
        arr = np.asarray(rows, float)
        arr = arr[np.argsort(arr[:, 0])]
        curves.append(_fit(SummaryCurve(name, arr[:, 0]*1e-6, arr[:, 1], arr[:, 2])))
    return curves, corrections


def is_summary_calibration(path):
    return os.path.splitext(path)[1].lower() in (".xlsx", ".xlsm")
