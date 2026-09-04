import tempfile
import unittest
from pathlib import Path

import numpy as np
from openpyxl import Workbook

from summary_calibration import load_summary_calibration
from unmix import _quantify_summary_map


def make_book(path, typo=False):
    wb = Workbook(); ws = wb.active
    for j, name in enumerate(("DQ", "TBZ", "THI")):
        col = 2 + 4*j
        ws.cell(1, col, name); ws.cell(2, col, "Mean"); ws.cell(2, col+1, "std")
        for row, c in enumerate((1.0, 10.0, 50.0, 100.0, 500.0), 3):
            entered = 9.0 if typo and name == "THI" and c == 10 else c
            response = 100 + 1000*(2e4*c*1e-6)/(1+2e4*c*1e-6)
            ws.cell(row, col-1, entered); ws.cell(row, col, response)
            ws.cell(row, col+1, 10.0)
    wb.save(path)


class SummaryCalibrationTest(unittest.TestCase):
    def test_typo_and_ratio_preservation(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp)/"cal.xlsx"; make_book(path, typo=True)
            curves, fixes = load_summary_calibration(path)
            self.assertEqual(fixes, ["THI: 9 -> 10 uM"])
            self.assertAlmostEqual(curves[2].concentration_M[1], 10e-6)
            ratios = np.array([[.2, .3, .5], [.1, .6, .3]])
            conc, *_ = _quantify_summary_map(
                str(path), ["DQ", "TBZ", "THI"], np.eye(3),
                np.array([[350., 350., 350.], [600., 600., 600.]]),
                np.array([1000., 1100., 1200.]), np.ones(2, bool), ratios,
                {"DQ": 1000, "TBZ": 1100, "THI": 1200})
            self.assertTrue(np.allclose(conc/conc.sum(1, keepdims=True), ratios))


if __name__ == "__main__":
    unittest.main()
